#!/usr/bin/env python3
"""Close run 2026-10-28: the 2026-10-23 refine logic (scripts/refine_2026_10_23.py, copied and extended) plus the D-batch part-B results,
the duplicate decisions (dupe-review-2026-10-28.csv), the mirror licence look-ups, and the close-run outputs:
inventory-refined-2026-10-28.csv, inventory.csv, browser-queue-2026-10-28.csv, notes-for-tim-2026-10-28.md, ticket-2026-10-28/, the workbook and close-report-2026-10-28.md.
Re-runnable after each batch. Run from docs/discovery:  python scripts/refine_2026_10_28.py
"""
import csv, json, os, re, sys, glob, collections, datetime
from collections import Counter, defaultdict, OrderedDict

D = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.join(D, "scripts"))
import verify_links as VL  # summarise(), CHECKS, domain_of, registered_domain

DATE = "2026-10-28"; PREV = "2026-10-23"; TODAY = "2026-10-28"; READ_DAY = {"B": "2026-10-22", "C": "2026-10-22", "m": "2026-10-22", "D": "2026-10-28"}
A_CHECKS = VL.CHECKS
B_CHECKS = ["chk_license", "chk_temporal", "chk_recency", "chk_cadence", "chk_version", "chk_unit", "chk_coverage", "chk_access", "chk_join_key", "chk_format_gaia"]
C_CHECKS = ["chk_in_gaia", "chk_recipe_health"]
ALL_CHECKS = A_CHECKS + B_CHECKS + C_CHECKS
CONFIRMED = ["license_confirmed", "license_class", "temporal_confirmed", "latest_year_confirmed", "update_cadence_confirmed", "version_confirmed",
             "finest_unit_confirmed", "coverage_confirmed", "access_confirmed", "join_key_confirmed", "formats_confirmed", "withdrawn_notice_B", "partB_pages_read", "partB_notes"]
WEIGHTS = {"chk_link_live": 20, "chk_soft404": 10, "chk_withdrawn": 10, "chk_license": 15, "chk_unit": 15, "chk_temporal": 7.5, "chk_recency": 7.5,
           "chk_access": 10, "chk_join_key": 10, "chk_machine_metadata": 5}
VAL = {"pass": 1.0, "warn": 0.5, "fail": 0.0}

# Part C: recipe health per gaiaCatalog folder, from docs/deployment-test/log-*.md and docs/readiness/ (see partC-recipe-evidence-2026-10-23.csv)
RECIPE = {
    "us_2023_county_tl": ("ok", "deployment-test/log-2026-10-07-run4.md row 16: ingest Pass, 3,235 rows"),
    "us_2014_2019_monthly_pm25_by_county_cdc": ("ok", "deployment-test/log-2026-10-07-run4.md rows 17, 21, 22: ingest, load, join Pass"),
    "synthetic_county_ses": ("ok", "deployment-test/log-2026-10-07-run4.md row 18: ingest Pass"),
    "us_2022_county_tl": ("ok", "deployment-test/log-2026-10-01-run2.md: ingested (tutorial run 2)"),
    "us_2022_annual_aqi_by_county": ("ok", "deployment-test/log-2026-10-01-run2.md: ingested (tutorial run 2)"),
    "us_2020_census_county_fips": ("ok", "deployment-test/log-2026-10-01-run2.md: ingested (tutorial run 2)"),
    "global_pm25_concentration_1998_2016": ("broken", "readiness/evidence/working-log-2026-09-23.md and deployment-test/log-2026-10-01-run2.md: SEDAC file not retrievable (SSL failure, then timeouts)"),
    "ma_2022_svi_tract": ("broken", "readiness/findings-2026-09-24.md: dated SVI variable failed to load (PostgREST 400); ingest itself ran"),
    "ma_2018_svi_tract": ("untested", "readiness/evidence/working-log-2026-09-23.md: mentioned, no ingest result recorded"),
    "ma_2020_svi_tract": ("untested", "readiness/findings-2026-09-24.md: mentioned, no ingest result recorded"),
    "us_2018_svi_county": ("untested", "deployment-test/log-2026-10-01-run2.md: mentioned, no ingest result recorded"),
    "us_2023_annual_aqi_by_county": ("untested", "readiness/findings-2026-09-24.md: mentioned, no ingest result recorded"),
}
# Step 3.1 (2026-10-28): licence look-up for the four mirrors. Pages read with WebFetch on 2026-10-28: landing page, About page, GitHub licence, publisher terms page.
MIRROR_LICENSE = {
    "us_epa_ejscreen_pedp": ("fail", "", "unknown", "no licence or terms on https://screening-tools.com/epa-ejscreen or /about (only 'This is an initial reconstruction of Version 2.3 of EPA's EJScreen'); the Harvard Dataverse record (doi:10.7910/DVN/RLR5AX) returns an empty page to scripts; EPA original is US Government work but the mirror does not say so",
                             "https://screening-tools.com/epa-ejscreen; https://screening-tools.com/about"),
    "us_ceq_cejst_pedp": ("fail", "", "unknown", "no licence on the downloads or About page ('This is an unofficial copy of the CEJST Tool'); the PEDP GitHub repo's LICENSE.md is AGPL for 'portions of this repo' (code) while the README badge still says CC0-1.0 (the upstream DOI-DO repo); neither statement covers the data",
                          "https://public-environmental-data-partners.github.io/j40-cejst-2/en/downloads; https://public-environmental-data-partners.github.io/j40-cejst-2/en/about; https://github.com/public-environmental-data-partners/j40-cejst-2; https://raw.githubusercontent.com/Public-Environmental-Data-Partners/j40-cejst-2/main/LICENSE.md"),
    "us_hifld_next_pedp": ("fail", "", "unknown", "no licence or terms on the landing or About page (only 'Use the data layers for your work: planning, research, or analysis'); data 'Managed by the Department of Homeland Security (DHS)', preserved by the Data Rescue Project; operators PEDP and Fulton Ring; contact only via a feedback survey",
                           "https://hifld.publicenvirodata.org; https://hifld.publicenvirodata.org/about"),
    "us_usgs_3dhp": ("pass", "public domain (USGS-produced data; credit requested: 'Credit: U.S. Geological Survey')", "open", "publisher-level statement, not on the 3DHP pages: USGS Copyrights and Credits says 'USGS-authored or produced data and information are considered to be in the U.S. Public Domain' and asks that proper credit be given; the 3DHP program and access pages state no licence themselves",
                     "https://www.usgs.gov/information-policies-and-instructions/copyrights-and-credits; https://www.usgs.gov/3d-hydrography-program; https://www.usgs.gov/3d-hydrography-program/access-3dhp-data-products"),
}
MIRROR_EXTRA = {  # other facts read on 2026-10-28 (3DHP access page)
    "us_usgs_3dhp": {"update_cadence_confirmed": "web services quarterly (EDH additions); downloadable product early each fiscal year with a DOI per annual version", "version_confirmed": "web services updated July 2026; 2025 downloadable product (ScienceBase 6930876ed4be027594fed169)",
                     "access_confirmed": "open_download", "chk_cadence": ("pass", "quarterly services, annual download", "https://www.usgs.gov/3d-hydrography-program/access-3dhp-data-products"),
                     "chk_version": ("warn", "release dates named (July 2026 services; 2025 download), no version number", "https://www.usgs.gov/3d-hydrography-program/access-3dhp-data-products"),
                     "chk_access": ("pass", "feature service https://3dhp.nationalmap.gov/arcgis/rest/services/usgs_3dhp_all/FeatureServer and ScienceBase download, no login named", "https://www.usgs.gov/3d-hydrography-program/access-3dhp-data-products")},
}
STRONG_WORDS = re.compile(r"(retired|discontinued|decommissioned|no longer (available|maintained|updated|supported)|has been archived|published for the last time)", re.I)


def read_csv(p):
    return list(csv.DictReader(open(p, encoding="utf-8"))) if os.path.exists(p) else []


def write_csv(p, rows, cols):
    with open(p, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore"); w.writeheader()
        for r in rows:
            w.writerow({c: ("" if r.get(c) is None else r.get(c)) for c in cols})


# ----------------------------------------------------------------------------- load
trim = read_csv(f"{D}/trim-2026-10-22.csv")
A_latest = {r["id"]: r for r in read_csv(f"{D}/verify-A-2026-10-21.csv")}
A_mirror = {r["id"]: r for r in read_csv(f"{D}/partB/mirrors-A/verify-A-{PREV}.csv")}
mirrors = read_csv(f"{D}/partB/mirrors-{PREV}.csv")
PREV_REFINED = {r["id"]: r for r in read_csv(f"{D}/_to_delete/inventory-refined-{PREV}-before-close.csv")}
DUPES = read_csv(f"{D}/dupe-review-{DATE}.csv")
QUEUE_CLOSE = [x[0] for x in json.load(open(f"{D}/partB/queue-close.json"))] if os.path.exists(f"{D}/partB/queue-close.json") else []
BROWSER = {r["id"]: r for r in read_csv(f"{D}/browser-sample-{DATE}.csv") if r.get("opens", "").startswith("yes") or r.get("opens", "").startswith("no")}
BROWSER_TIER = {  # from browser-sample-2026-10-28.md, "Suggested status changes" (Robert, 2026-10-28)
    "us_ahrq_sdoh_database": "Silver", "us_brandeis_coi": "Silver", "us_hud_chas": "Gold",
    "us_hud_fair_market_rents": "Silver", "us_eviction_lab": "Silver", "us_ihme_county_mortality_life_expectancy": "Silver", "br_ipea_ivs": "Silver", "us_cdc_covid_community_levels": "Silver",
    "us_cms_hospital_general_information": "Silver", "us_cms_nursing_home_provider_info": "Silver", "us_hifld_hospitals": "Silver", "us_epa_campd_power_plant_emissions": "Silver",
    "global_sedac_pm25_countries_urban_1998_2016": "Silver",
    "us_faa_airport_noise_exposure_maps": "Trimmed", "hdx_kenya_number_of_hospital_icu_beds_per_county": "Trimmed", "us_enverus_well_data": "Trimmed",
    "us_cdc_heat_health_tracker": "Bronze",
    "us_sedac_amini_pm25_components_50m": "Failed", "us_sedac_di_pm25_1km_2000_2016": "Failed", "eu_elapse_lur_surfaces": "Failed", "us_hud_public_housing_developments": "Failed",
    "us_nrdc_lead_service_lines": "Browser needed",
}
BROWSER_NOTE = f"browser check {DATE}"
ACCESS_MAP = lambda a: "open_download" if a.startswith("download") or a.startswith("custom query") else "restricted" if a in ("subscription",) else "" if a in ("dashboard only", "") else "open_download"
TLS = json.load(open(f"{D}/partB/tls-diagnosis-2026-10-22.json"))
link = read_csv(f"{D}/gaia-link-2026-10-20.csv")
log = read_csv(f"{D}/verify-log.csv")
runs = sorted({r["run_id"] for r in log})
RUN_A, RUN_1 = runs[0], (runs[-1] if len(runs) > 1 else "")

partB = {}
for f in sorted(glob.glob(f"{D}/partB/results/*.json")):
    try:
        j = json.load(open(f, encoding="utf-8"))
    except Exception as e:
        print("WARNING: bad json", f, e, file=sys.stderr); continue
    for r in j.get("rows", []):
        partB[r["id"]] = (r, os.path.basename(f))
print(f"part B results for {len(partB)} rows from {len(glob.glob(f'{D}/partB/results/*.json'))} files", file=sys.stderr)

A_FIELDS = [c for c in A_CHECKS] + [c + "_detail" for c in A_CHECKS] + ["http_status", "final_url", "response_ms", "content_type", "page_title", "checked_at",
            "access_http_status", "access_final_url", "access_content_type", "metadata_kind", "persistent_id", "wayback_url", "wayback_date", "url_health"]

rows = []
for r in trim:
    r = dict(r); r["mirror_of"] = ""
    rows.append(r)
for m in mirrors:
    r = {k: "" for k in trim[0].keys()}; r.update(m); r["keep"] = "Y"; r["trim_reason"] = ""
    r["latest_year"] = ""
    rows.append(r)
# step 2 (2026-10-28): apply the duplicate decisions
DUPE_DROP, DUPE_NOTE, DUPE_ACCESS = {}, defaultdict(list), {}
for d in DUPES:
    ids = d["ids"].split(";"); prim = ids[0]
    if d["decision"] in ("same product", "host copy"):
        for alt in ids[1:]:
            DUPE_DROP[alt] = f"{'duplicate' if d['decision'] == 'same product' else 'host copy'} of {prim}"
        for part in d["applied"].split(";"):
            part = part.strip()
            if part.startswith("access_url on primary:"):
                DUPE_ACCESS[prim] = part.split(":", 1)[1].strip()
            elif part.startswith("note on primary:") or part.startswith("note:"):
                DUPE_NOTE[prim].append(part.split(":", 1)[1].strip())
    DUPE_NOTE[prim].append(f"dup review {DATE} ({d['group']}): {d['decision']}")
for r in rows:
    b = BROWSER.get(r["id"])
    if b and BROWSER_TIER.get(r["id"]) == "Trimmed" and r["keep"] == "Y":
        r["keep"] = "N"; r["trim_reason"] = f"{BROWSER_NOTE}: " + b["notes"][:160]
    if r["id"] == "us_ahrq_sdoh_database":
        r["site_url"] = "https://www.ahrq.gov/data/innovations/clh-data.html"
        r["dataset_name"] = "AHRQ Community-Level Health (CLH) Database (formerly the SDOH Database)"
    if r["id"] in DUPE_DROP and r["keep"] == "Y":
        r["keep"] = "N"; r["trim_reason"] = DUPE_DROP[r["id"]]
    if r["id"] in DUPE_ACCESS and not r.get("access_url"):
        r["access_url"] = DUPE_ACCESS[r["id"]]
    if r["id"] in DUPE_NOTE:
        r["notes"] = ((r.get("notes") or "").rstrip("; ") + "; " if r.get("notes") else "") + "; ".join(DUPE_NOTE[r["id"]])
DUPE_DROP_KEPT = {i: v for i, v in DUPE_DROP.items() if i in PREV_REFINED}  # alternates that were actually kept rows
K = [r for r in rows if r["keep"] == "Y"]
T = [r for r in rows if r["keep"] != "Y"]

# step-1 changes (what the re-check recovered)
step1 = []
if RUN_1:
    cur = {(x["row_id"], x["check"]): x for x in log if x["run_id"] == RUN_1}
    old = {(x["row_id"], x["check"]): x for x in log if x["run_id"] == RUN_A}
    for k, x in cur.items():
        if k[1] == "chk_link_live":
            step1.append((k[0], old.get(k, {}).get("result", ""), x["result"], x["detail"][:80]))

# ----------------------------------------------------------------------------- per-row merge
for r in K:
    a = A_mirror.get(r["id"]) or A_latest.get(r["id"])
    if a:
        for c in A_FIELDS:
            r[c] = a.get(c, r.get(c, ""))
    # TLS diagnosis
    if r["id"] in TLS:
        kind, txt = TLS[r["id"]]
        r["chk_tls_detail"] = f"{kind}: {txt}"
        r["chk_link_live_detail"] = f"tls_error ({kind}; not bypassed)"
    # mirror hand judgement: the EJScreen mirror page quotes the EPA discontinuation; the mirror itself is live
    if r["id"] == "us_epa_ejscreen_pedp":
        r["chk_withdrawn"] = "pass"; r["chk_withdrawn_detail"] = "withdrawal sentence on the page refers to the EPA original (discontinued 2025-02-05); the PEDP mirror is live (hand judgement, refine 2026-10-23)"
    # recompute url_health from the (possibly updated) A checks
    tmp = {c: r.get(c, "") for c in A_CHECKS}; tmp.update({c + "_detail": r.get(c + "_detail", "") for c in A_CHECKS})
    VL.summarise(tmp); r["url_health"] = tmp["url_health"]
    if r["id"] in TLS:
        r["url_health"] = "blocked"  # certificate problems go to the browser queue (refine prompt, step 1)
    # part B
    pb = partB.get(r["id"])
    for c in CONFIRMED:
        r[c] = ""
    r["partB_blocked"] = ""
    if pb:
        b, src = pb
        r["partB_source"] = src
        r["partB_blocked"] = "TRUE" if b.get("blocked") else "FALSE"
        for c in B_CHECKS:
            x = (b.get("checks") or {}).get(c) or {}
            r[c] = x.get("result") or "not_checked"; r[c + "_detail"] = x.get("detail", ""); r[c + "_evidence"] = x.get("evidence", "")
        r["license_confirmed"] = b.get("license_confirmed", ""); r["license_class"] = b.get("license_class", "")
        r["temporal_confirmed"] = b.get("temporal_confirmed", ""); r["latest_year_confirmed"] = b.get("latest_year") if b.get("latest_year") is not None else ""
        r["update_cadence_confirmed"] = b.get("update_cadence_confirmed", ""); r["version_confirmed"] = b.get("version_confirmed", "")
        r["finest_unit_confirmed"] = b.get("finest_unit_confirmed", ""); r["coverage_confirmed"] = b.get("coverage_confirmed", "")
        r["access_confirmed"] = b.get("access_confirmed", ""); r["join_key_confirmed"] = b.get("join_key_confirmed", ""); r["formats_confirmed"] = b.get("formats_confirmed", "")
        r["withdrawn_notice_B"] = b.get("withdrawn_notice", ""); r["partB_pages_read"] = "; ".join(b.get("pages_read") or []); r["partB_notes"] = b.get("notes", "")
        if b.get("blocked"):
            r["url_health"] = "blocked"
            if "blocked" not in r["chk_link_live_detail"]:
                r["chk_link_live_detail"] += " | content step: fetcher got a challenge/empty page (blocked)"
    else:
        r["partB_source"] = ""
        for c in B_CHECKS:
            r[c] = "not_checked"; r[c + "_detail"] = "part B not run on this row"; r[c + "_evidence"] = ""
    if r["id"] in MIRROR_LICENSE:
        res, lic, cls, det, ev_ = MIRROR_LICENSE[r["id"]]
        r["chk_license"] = res; r["chk_license_detail"] = det; r["chk_license_evidence"] = ev_
        r["license_confirmed"] = lic; r["license_class"] = cls
        r["partB_pages_read"] = "; ".join(OrderedDict.fromkeys([p for p in (r.get("partB_pages_read", "").split("; ") + ev_.split("; ")) if p]))
        for k2, v2 in MIRROR_EXTRA.get(r["id"], {}).items():
            if isinstance(v2, tuple):
                r[k2], r[k2 + "_detail"], r[k2 + "_evidence"] = v2
            else:
                r[k2] = v2
    # browser sample (2026-10-28): a person opened the page; fields and tier from browser-sample-2026-10-28.csv / .md
    b = BROWSER.get(r["id"])
    if b and r["id"] in BROWSER_TIER:
        tgt = BROWSER_TIER[r["id"]]
        r["notes"] = ((r.get("notes") or "").rstrip("; ") + "; " if r.get("notes") else "") + f"{BROWSER_NOTE}: " + b["notes"][:200]
        if b["opens"].startswith("yes"):
            r["url_health"] = "healthy"; r["chk_link_live"] = "pass"; r["chk_link_live_detail"] = f"opened in a browser ({BROWSER_NOTE})"
            r["chk_soft404"] = "pass"; r["chk_soft404_detail"] = f"page names the dataset ({BROWSER_NOTE})"
            if r["chk_withdrawn"] in ("not_checked", "n.a.", ""):
                r["chk_withdrawn"] = "pass"; r["chk_withdrawn_detail"] = f"no retirement notice seen ({BROWSER_NOTE})"
            lic = b["license"].strip()
            if lic and lic not in ("n/a", "not stated", "not seen"):
                r["license_confirmed"] = lic
                r["license_class"] = "open" if lic.startswith("CC BY") else "unknown" if "assumed" in lic else "non-commercial" if "non-commercial" in lic else "restricted" if "commercial" in lic else "attribution" if "cite" in lic else "unknown"
                r["chk_license"] = "pass" if r["license_class"] in ("open", "attribution") else "warn" if "assumed" in lic else "warn" if r["license_class"] == "non-commercial" else "fail"
                r["chk_license_detail"] = f"{lic} ({BROWSER_NOTE})" + ("; publisher status assumed from the page, not stated" if "assumed" in lic else ""); r["chk_license_evidence"] = b["site_url"]
            elif lic:
                r["chk_license"] = "fail"; r["chk_license_detail"] = f"no licence seen on the landing page ({BROWSER_NOTE})"; r["chk_license_evidence"] = b["site_url"]
            if b["years"].strip() and b["years"] not in ("varies", "not seen on landing page"):
                r["temporal_confirmed"] = b["years"]; r["chk_temporal"] = "pass"; r["chk_temporal_detail"] = f"{b['years']} ({BROWSER_NOTE})"; r["chk_temporal_evidence"] = b["site_url"]
            if b["latest_year"].strip().isdigit():
                r["latest_year_confirmed"] = b["latest_year"]; y = int(b["latest_year"])
                r["chk_recency"] = "pass" if y >= 2020 else "warn" if y >= 2000 else "fail"; r["chk_recency_detail"] = f"latest year {y} ({BROWSER_NOTE})"; r["chk_recency_evidence"] = b["site_url"]
            if b["finest_unit"].strip():
                r["finest_unit_confirmed"] = b["finest_unit"]; r["chk_unit"] = "pass" if b["finest_unit"].split(" ")[0].lower().startswith(r["finest_unit"][:4].lower()) else "warn"
                r["chk_unit_detail"] = f"{b['finest_unit']} ({BROWSER_NOTE}; listed: {r['finest_unit']})"; r["chk_unit_evidence"] = b["site_url"]
            acc = ACCESS_MAP(b["access"].strip().lower())
            if acc:
                r["access_confirmed"] = acc; r["chk_access"] = "pass" if acc == "open_download" else "fail"; r["chk_access_detail"] = f"{b['access']} ({BROWSER_NOTE})"; r["chk_access_evidence"] = b["site_url"]
            if b["formats"].strip() and b["formats"] not in ("none",):
                r["formats_confirmed"] = b["formats"]; r["chk_format_gaia"] = "warn" if b["formats"] in ("PDF",) else "pass"; r["chk_format_gaia_detail"] = f"{b['formats']} ({BROWSER_NOTE})"; r["chk_format_gaia_evidence"] = b["site_url"]
            if "iscontinued" in b["notes"]:
                r["withdrawn_notice_B"] = r.get("withdrawn_notice_B") or b["notes"][:120]
        elif tgt == "Failed":
            r["url_health"] = "dead"; r["chk_link_live"] = "fail"; r["chk_link_live_detail"] = f"{b['opens']} ({BROWSER_NOTE}): " + b["notes"][:120]
        r["browser_tier"] = tgt
    else:
        r["browser_tier"] = ""
    # part C
    r["chk_in_gaia"] = "pass" if r.get("in_gaia", "").strip() else "n.a."
    r["chk_in_gaia_detail"] = ("gaiaCatalog folder(s): " + r["in_gaia"]) if r.get("in_gaia", "").strip() else "not in gaiaCatalog"
    if r.get("in_gaia", "").strip():
        folders = [f.strip() for f in re.split(r"[;,]", r["in_gaia"]) if f.strip()]
        res = [RECIPE.get(f, ("untested", "no result in docs/deployment-test/log-*.md or docs/readiness/")) for f in folders]
        if any(x[0] == "broken" for x in res):
            r["chk_recipe_health"] = "fail"
        elif any(x[0] == "ok" for x in res):
            r["chk_recipe_health"] = "pass"
        else:
            r["chk_recipe_health"] = "not_checked"
        r["chk_recipe_health_detail"] = "; ".join(f"{f}: {x[0]} ({x[1]})" for f, x in zip(folders, res))[:400]
        r["recipe_health"] = "; ".join(f"{f}: {x[0]}" for f, x in zip(folders, res))
    else:
        r["chk_recipe_health"] = "n.a."; r["chk_recipe_health_detail"] = "not in gaiaCatalog"; r["recipe_health"] = ""
    # latest year: part B first, then the listing
    r["latest_year_final"] = str(r["latest_year_confirmed"] or r.get("latest_year", "") or "")
    # summary
    app = [c for c in ALL_CHECKS if r.get(c) in VAL]
    r["checks_applicable"] = len(app); r["checks_passed"] = sum(1 for c in app if r[c] == "pass")
    r["checks_warned"] = sum(1 for c in app if r[c] == "warn"); r["checks_failed"] = sum(1 for c in app if r[c] == "fail")
    wsum = sum(WEIGHTS[c] for c in app if c in WEIGHTS)
    r["health_score"] = round(100 * sum(WEIGHTS[c] * VAL[r[c]] for c in app if c in WEIGHTS) / wsum) if wsum else ""
    r["all_passing"] = "TRUE" if app and r["checks_passed"] == len(app) else "FALSE"
    r["verify_blockers"] = ";".join([c[4:] for c in ALL_CHECKS if r.get(c) == "fail"] + [c[4:] + "(warn)" for c in ALL_CHECKS if r.get(c) == "warn"])
    dates = [r.get("checked_at", "")[:10]] + ([READ_DAY.get(r["partB_source"][:1], TODAY)] if pb else []) + ([TODAY] if r["id"] in MIRROR_LICENSE or r["id"] in BROWSER_TIER else [])
    r["verified_on"] = max(d for d in dates if d) if any(dates) else ""
    B_done = any(r[c] in VAL for c in B_CHECKS)
    ll = r["chk_link_live"]
    if r["url_health"] == "blocked":
        tier = "Browser needed"
    elif ll == "fail" or r["chk_withdrawn"] == "fail" or r["chk_soft404"] == "fail" or r["chk_recency"] == "fail":
        tier = "Failed"
    elif all(r[c] == "pass" for c in ("chk_link_live", "chk_soft404", "chk_withdrawn", "chk_license", "chk_unit", "chk_temporal", "chk_access")) and r["chk_recency"] in ("pass", "warn"):
        tier = "Gold"
    elif ll == "pass" and B_done and not any(r[c] == "fail" for c in ("chk_license", "chk_unit", "chk_temporal", "chk_recency")):
        tier = "Silver"
    elif ll in ("pass", "warn") or ll == "n.a.":
        tier = "Bronze"
    else:
        tier = "Failed"
    if r.get("browser_tier"):
        tier = r["browser_tier"]  # suggested status (browser-sample-2026-10-28.md); the checks above record what the page showed
    r["verify_tier"] = tier
    # public access
    acc = (r.get("access_confirmed") or r.get("access_type") or "").lower()
    r["public_access"] = "open" if acc in ("open_download", "api") else "registration" if acc in ("registration", "api_key") else "restricted" if acc in ("request", "restricted") else "unknown"

# ----------------------------------------------------------------------------- 1. refined CSV
RAW = ["dataset_name", "organization", "sub_organization", "theme", "variables_examples", "geography_units", "coverage", "temporal", "versions", "format",
       "access_type", "license", "update_cadence", "source_of_listing", "evidence_url", "confidence", "last_checked", "notes", "in_gaia", "record_type",
       "health_relevance", "finest_unit", "gaia_fit", "score_note", "unit_basis", "unit_confidence", "unit_evidence"]
COLS = ["id", "mirror_of", "tier", "priority", "site_url", "access_url", "verify_tier", "browser_tier", "health_score", "url_health", "all_passing", "verify_blockers", "verified_on",
        "checks_applicable", "checks_passed", "checks_warned", "checks_failed", "latest_year_final", "public_access", "recipe_health"]
for c in A_CHECKS:
    COLS += [c, c + "_detail"]
for c in B_CHECKS:
    COLS += [c, c + "_detail", c + "_evidence"]
for c in C_CHECKS:
    COLS += [c, c + "_detail"]
COLS += CONFIRMED + ["partB_blocked", "partB_source"] + ["http_status", "final_url", "response_ms", "content_type", "page_title", "checked_at", "access_http_status",
                                                        "access_final_url", "access_content_type", "metadata_kind", "persistent_id", "wayback_url", "wayback_date", "latest_year"] + RAW
TIER_ORDER = {"Gold": 0, "Silver": 1, "Bronze": 2, "Browser needed": 3, "Failed": 4}
hrn = lambda r: int(r["health_relevance"]) if str(r.get("health_relevance", "")).isdigit() else -1
K.sort(key=lambda r: (TIER_ORDER[r["verify_tier"]], r["theme"], -hrn(r), r["dataset_name"].lower()))
write_csv(f"{D}/inventory-refined-{DATE}.csv", K, COLS)

# ----------------------------------------------------------------------------- 2. promote Gold and Silver into inventory.csv
INV_COLS = ["id", "title", "publisher", "domain", "landing_url", "access", "geography", "coverage", "temporal", "format", "license", "update_cadence", "in_gaia", "recipe_health", "last_checked", "evidence", "notes"]
promoted = []
for r in K:
    if r["verify_tier"] in ("Gold", "Silver"):
        ev = "; ".join(OrderedDict.fromkeys(x for x in [r.get("chk_license_evidence"), r.get("chk_temporal_evidence"), r.get("chk_unit_evidence"), r.get("evidence_url")] if x))
        promoted.append({"id": r["id"], "title": r["dataset_name"], "publisher": r["organization"], "domain": r["theme"], "landing_url": r["site_url"],
                         "access": r.get("access_confirmed") or r["access_type"], "geography": r.get("finest_unit_confirmed") or r["geography_units"] or r["finest_unit"],
                         "coverage": r.get("coverage_confirmed") or r["coverage"], "temporal": r.get("temporal_confirmed") or r["temporal"],
                         "format": r.get("formats_confirmed") or r["format"], "license": r.get("license_confirmed") or r["license"],
                         "update_cadence": r.get("update_cadence_confirmed") or r["update_cadence"], "in_gaia": r["in_gaia"], "recipe_health": r["recipe_health"],
                         "last_checked": r["verified_on"], "evidence": ev,
                         "notes": f"verify_tier {r['verify_tier']}; health_score {r['health_score']}; " + (f"mirror of {r['mirror_of']}; " if r["mirror_of"] else "") + (r["notes"] or "")[:300]})
write_csv(f"{D}/inventory.csv", promoted, INV_COLS)

# ----------------------------------------------------------------------------- 3. browser queue, grouped by domain
BQ = []
for r in K:
    if r["url_health"] == "blocked" or r["id"] in TLS:
        d = r["chk_link_live_detail"]
        reason = ("certificate: " + TLS[r["id"]][0]) if r["id"] in TLS else "blocked_robots" if "blocked_robots" in d else "timeout" if "timeout" in d \
            else "connection error" if d.startswith("connection") else "fetcher got a challenge/empty page" if r["partB_blocked"] == "TRUE" and "blocked_bot" not in d else "blocked_bot (401/403/202 challenge)"
        todo = ["page loads and names this dataset", "not retired or discontinued"]
        if r["id"] in TLS:
            todo.insert(0, "browser accepts the certificate (" + TLS[r["id"]][0] + ")")
        for c, what in (("chk_license", "licence / terms"), ("chk_temporal", "years covered and latest year"), ("chk_unit", "finest spatial unit (listed: " + (r["finest_unit"] or "?") + ")"),
                        ("chk_access", "access route"), ("chk_format_gaia", "download formats")):
            if r[c] not in ("pass",):
                todo.append(what)
        BQ.append({"domain": VL.registered_domain(VL.domain_of(r["site_url"])) or "(no url)", "id": r["id"], "dataset_name": r["dataset_name"], "organization": r["organization"],
                   "theme": r["theme"], "priority": r["priority"], "health_relevance": r["health_relevance"], "site_url": r["site_url"], "access_url": r["access_url"],
                   "reason": reason, "http_status": r["http_status"], "what_to_confirm": "; ".join(todo), "archive_copy": r["wayback_url"]})
BQ.sort(key=lambda x: (x["domain"], x["priority"] == "", x["dataset_name"].lower()))
BQ_COLS = ["domain", "id", "dataset_name", "organization", "theme", "priority", "health_relevance", "site_url", "access_url", "reason", "http_status", "what_to_confirm", "archive_copy"]
write_csv(f"{D}/browser-queue-{DATE}.csv", BQ, BQ_COLS)

# ----------------------------------------------------------------------------- 4. notes for Tim (gaiaCatalog sources with link problems)
folder_of = defaultdict(list)
for l in link:
    folder_of[l["matched_id"]].append(l)
tim = []
for r in rows:  # all rows, kept or trimmed, that are in gaiaCatalog
    if not r.get("in_gaia", "").strip():
        continue
    a = A_latest.get(r["id"], r)
    uh = r.get("url_health") if r["keep"] == "Y" else a.get("url_health", "")
    wd = (r.get("chk_withdrawn") if r["keep"] == "Y" else a.get("chk_withdrawn", "")) == "fail"
    if uh in ("dead", "blocked") or wd:
        folders = [f.strip() for f in re.split(r"[;,]", r["in_gaia"]) if f.strip()]
        urls = {l["folder"]: l["catalog_url"] for l in folder_of.get(r["id"], [])}
        det = a.get("chk_link_live_detail", "") if not wd else "withdrawn: " + a.get("chk_withdrawn_detail", "")[:160]
        tim.append((folders, r, uh if not wd else "withdrawn", det, urls))
mirror_for = {m["mirror_of"]: m for m in mirrors}
SUGGEST = {
    "us_nuclear_power_plants": "HIFLD Next (PEDP archive of HIFLD Open) power-plant layers: https://hifld.publicenvirodata.org/collections (us_hifld_next_pedp), or the Data Rescue Project copy of HIFLD Open power plants (us_hifld_open_power_plants, kept, Bronze); the EIA Atlas nuclear page itself is 404",
    "global_pm25_concentration_1998_2016": "SEDAC moved to NASA Earthdata: catalog record https://www.earthdata.nasa.gov/data/catalog (search 'Annual PM2.5 Concentrations for Countries and Urban Areas'); the Harvard Dataverse copy (doi:10.7910/DVN/RHIFKL) is behind a 202 challenge for scripts",
}
L = [f"# Notes for Tim: gaiaCatalog sources with link problems ({DATE}, draft, not sent; supersedes the 2026-10-23 draft)", "",
     "**This assumes:** the scripted link checks of 2026-10-21 (part A) and the re-check of 2026-10-22; the page reads of 2026-10-22 and 2026-10-28 (part B); 'blocked' means the site refuses scripts, not that the data are gone. Catalog folder names are from `gaia-link-2026-10-20.csv`.", "",
     "| catalog folder | source URL checked | result (2026-10-22) | detail | suggested replacement |", "|---|---|---|---|---|"]
seen = set()
for folders, r, uh, det, urls in sorted(tim, key=lambda t: t[0][0]):
    for f in folders:
        if f in seen:
            continue
        seen.add(f)
        sug = SUGGEST.get(f, "")
        if not sug and r["id"] in mirror_for:
            m = mirror_for[r["id"]]; sug = f"{m['dataset_name']} ({m['id']}): {m['site_url']}"
        if not sug and uh == "blocked":
            sug = "none needed: opens in a browser; the recipe's download URL may still work"
        L.append(f"| `{f}` | {urls.get(f) or r['site_url']} | {uh} | {det.replace('|', '/')[:140]} | {sug or 'none found'} |")
L += ["", f"{len(seen)} folders listed. Recipe health from the deployment logs: " + "; ".join(f"`{k}` {v[0]}" for k, v in RECIPE.items()) + ".", "",
      "Also worth a look: the EJScreen and CEJST sources are withdrawn by their publishers; PEDP mirrors are in the refined inventory (`us_epa_ejscreen_pedp`, `us_ceq_cejst_pedp`), and USGS says 3DHP replaces the NHD (`us_usgs_3dhp`). None of the three PEDP pages states a licence (checked 2026-10-28); USGS data are public domain by USGS policy."]
L += ["", f"Browser sample ({DATE}): the Harvard Dataverse copy of the SEDAC PM2.5 countries/urban data (`global_sedac_pm25_countries_urban_1998_2016`, a gaiaCatalog source) opens fine in a browser under CC BY 4.0, so the scripted block is bot-protection, not a dead link. The SEDAC host itself (sedac.ciesin.columbia.edu) does not load in a browser either, and a CIESIN page is 404: check NASA Earthdata for any recipe still pointing at sedac.ciesin.columbia.edu."]
# 2026-10-28: what the close-run page reads (D batches) found about catalog sources
L += ["", f"## What the {DATE} page reads found about catalog sources", "",
      "Rows with a gaiaCatalog folder that were read for content in the close run (D batches), with the agent's findings. 'blocked' = the fetcher got a JavaScript shell or challenge page; open it in a browser.", "",
      "| catalog folder(s) | inventory row | result | licence | unit | years | note |", "|---|---|---|---|---|---|---|"]
for r in sorted(K, key=lambda r: r["in_gaia"]):
    if r.get("in_gaia", "").strip() and r["partB_source"].startswith("D"):
        L.append(f"| {r['in_gaia'].replace(';', '; ')} | `{r['id']}` | {'blocked' if r['partB_blocked'] == 'TRUE' else r['verify_tier']} | {(r['license_confirmed'] or r['chk_license'])[:60]} | {r['finest_unit_confirmed'] or '-'} | {(r['temporal_confirmed'] or '-')[:40]} | {(r['partB_notes'] or '').replace('|', '/')[:160]} |")
open(f"{D}/notes-for-tim-{DATE}.md", "w", encoding="utf-8").write("\n".join(L) + "\n")

# ----------------------------------------------------------------------------- 7. ticket package
TK = f"{D}/ticket-{DATE}"; os.makedirs(TK, exist_ok=True)
GF = {"now": "Joins today", "needs_aggregation": "Needs a count/distance step", "needs_raster": "Needs raster support", "not_linkable": "Not linkable as listed", "unknown": "Unknown"}
UH = {"healthy": "Live", "degraded": "Live, page needs a look", "moved": "Live (moved)", "blocked": "Site blocks scripts (check in browser)", "dead": "Broken link / withdrawn", "n.a.": "No URL"}


def trow(r):
    return {"status": r["verify_tier"], "dataset": r["dataset_name"], "publisher": r["organization"], "theme": r["theme"], "coverage": r.get("coverage_confirmed") or r["coverage"],
            "finest_unit": r.get("finest_unit_confirmed") or r["finest_unit"], "years": r.get("temporal_confirmed") or r["temporal"], "latest_year": r["latest_year_final"],
            "gaia_fit": GF.get(r["gaia_fit"], r["gaia_fit"]), "health_relevance": r["health_relevance"], "link_status": UH.get(r["url_health"], r["url_health"]),
            "in_gaiacatalog": r["in_gaia"], "format": r.get("formats_confirmed") or r["format"], "access": r.get("access_confirmed") or r["access_type"], "public_access": r["public_access"],
            "license_as_listed": r["license"], "license_confirmed": r["license_confirmed"], "license_class": r["license_class"] or ("" if r["partB_source"] else "not checked"),
            "finest_unit_confirmed": r["finest_unit_confirmed"], "join_key_confirmed": r["join_key_confirmed"],
            "evidence": "; ".join(OrderedDict.fromkeys(x for x in [r.get(c + "_evidence") for c in B_CHECKS] if x)),
            "health_score": r["health_score"], "verified_on": r["verified_on"], "landing_page": r["site_url"], "archive_copy": r["wayback_url"], "id": r["id"], "mirror_of": r["mirror_of"]}


def what_to_check(r):
    w = []
    if r["verify_tier"] == "Browser needed":
        w.append("open in a browser (site blocks scripts" + (", certificate problem" if r["id"] in TLS else "") + ")")
    if r["gaia_fit"] == "needs_raster": w.append("raster support in Gaia")
    if r["gaia_fit"] == "needs_aggregation": w.append("count/distance step for points or lines")
    if r["gaia_fit"] == "unknown": w.append("confirm the spatial unit")
    if r["partB_source"] == "": w.append("content not yet read (licence, years, unit)")
    for c, what in (("chk_license", "licence"), ("chk_unit", "unit"), ("chk_temporal", "years"), ("chk_access", "access"), ("chk_join_key", "join key"), ("chk_format_gaia", "loadable format")):
        if r[c] == "warn": w.append(f"{what}: " + (r[c + "_detail"] or "")[:60])
        elif r[c] == "fail": w.append(f"{what} FAILED: " + (r[c + "_detail"] or "")[:60])
    if r["chk_soft404"] == "warn": w.append("page match weak: " + r["chk_soft404_detail"][:50])
    return "; ".join(OrderedDict.fromkeys(w))


TCOLS = list(trow(K[0]).keys())
ref = [r for r in K if r["health_relevance"] not in ("2", "3") and r["verify_tier"] != "Failed"]
ver = [r for r in K if r["health_relevance"] in ("2", "3") and r["verify_tier"] in ("Gold", "Silver")]
inv = [r for r in K if r["health_relevance"] in ("2", "3") and r["verify_tier"] in ("Bronze", "Browser needed")]
fail = [r for r in K if r["verify_tier"] == "Failed"]
write_csv(f"{TK}/1-verified.csv", [trow(r) for r in ver], TCOLS)
write_csv(f"{TK}/2-to-investigate.csv", [dict(trow(r), what_to_check=what_to_check(r)) for r in inv], TCOLS + ["what_to_check"])
write_csv(f"{TK}/3-reference-layers.csv", [trow(r) for r in ref], TCOLS)
why = lambda r: r["trim_reason"] if r["keep"] != "Y" else "FAILED verification: " + ("; ".join(f"{c[4:]}: {r[c+'_detail'][:80]}" for c in ("chk_link_live", "chk_withdrawn", "chk_soft404", "chk_recency") if r[c] == "fail"))
write_csv(f"{TK}/4-trimmed.csv", [{"dataset": r["dataset_name"], "publisher": r["organization"], "why_trimmed": why(r), "latest_year": r.get("latest_year_final") or r.get("latest_year", ""),
                                   "landing_page": r["site_url"], "id": r["id"]} for r in T + fail], ["dataset", "publisher", "why_trimmed", "latest_year", "landing_page", "id"])

# ISSUE update (what changed since the 2026-10-23 package)
old = {k: read_csv(f"{D}/ticket-{PREV}/{k}.csv") for k in ("1-verified", "2-to-investigate", "3-reference-layers", "4-trimmed")}
old_ids = {k: {x.get("id") for x in v} for k, v in old.items()}
new_ids = {"1": {r["id"] for r in ver}, "2": {r["id"] for r in inv}, "3": {r["id"] for r in ref}, "4": {r["id"] for r in T + fail}}
moved = Counter()
for r in K + [x for x in T if x["id"] in DUPE_DROP]:
    was = "1" if r["id"] in old_ids["1-verified"] else "2" if r["id"] in old_ids["2-to-investigate"] else "3" if r["id"] in old_ids["3-reference-layers"] else "4" if r["id"] in old_ids["4-trimmed"] else "new"
    now = "1" if r["id"] in new_ids["1"] else "2" if r["id"] in new_ids["2"] else "3" if r["id"] in new_ids["3"] else "4"
    if was != now:
        moved[(was, now)] += 1
tiers = Counter(r["verify_tier"] for r in K)
tiers_prev = Counter(PREV_REFINED[i]["verify_tier"] for i in PREV_REFINED)
tier_moves = Counter((PREV_REFINED[r["id"]]["verify_tier"], r["verify_tier"]) for r in K if r["id"] in PREV_REFINED and PREV_REFINED[r["id"]]["verify_tier"] != r["verify_tier"])
newly_read = [r for r in K if r["partB_source"].startswith("D")]
newly_read_ok = [r for r in newly_read if r["partB_blocked"] != "TRUE"]
newly_blocked = [r for r in newly_read if r["partB_blocked"] == "TRUE"]
lic = Counter(r["license_class"] for r in K if r["license_class"])
dup_dec = Counter(d["decision"] for d in DUPES)
reg = [r for r in K if r["public_access"] in ("registration", "restricted") and r["verify_tier"] in ("Gold", "Silver", "Bronze")]
IS = [f"## Update ({DATE}): first-round assessment closed", "",
      f"**What changed since the {PREV} post.** The remaining live, unread datasets ({len(newly_read)} rows: health relevance 3 with an unknown or non-linkable unit, relevance 2 needing a count/distance step or raster support, reference layers, and the rest) went through the same content check as the shortlist: an agent read the landing page and one documentation page and confirmed licence, years, unit, coverage, access, join key and formats with evidence URLs. {len(newly_blocked)} of them turned out to be JavaScript-only or challenge pages and join the browser queue. Every flagged possible-duplicate group among the kept rows was decided, and the four mirrors' licences were looked up. Every kept dataset has a verification tier; {sum(1 for r in K if r['partB_source'])} of {len(K)} have been read for content.", "",
      "| Tier | Meaning | Was | Now |", "|---|---|---|---|",
      f"| Gold | link, page, licence, unit, years and access all confirmed; data runs to 2000 or later | {tiers_prev.get('Gold', 0)} | {tiers.get('Gold', 0)} |",
      f"| Silver | link confirmed; content read with no failure on licence, unit, years or recency | {tiers_prev.get('Silver', 0)} | {tiers.get('Silver', 0)} |",
      f"| Bronze | link live; content read but incomplete, or not readable by script | {tiers_prev.get('Bronze', 0)} | {tiers.get('Bronze', 0)} |",
      f"| Browser needed | site blocks scripts, JavaScript-only page or certificate problem; a person must open it | {tiers_prev.get('Browser needed', 0)} | {tiers.get('Browser needed', 0)} |",
      f"| Failed | link dead, page generic, publisher says withdrawn, or data older than 2000 | {tiers_prev.get('Failed', 0)} | {tiers.get('Failed', 0)} |", "",
      "Tier movements: " + (", ".join(f"{a} → {b}: {n}" for (a, b), n in sorted(tier_moves.items(), key=lambda x: -x[1])) or "none") + ".", "",
      "### Files (same four as before)", "", "| File | Rows | Was | What it is |", "|---|---|---|---|",
      f"| `1-verified.csv` | {len(ver)} | {len(old['1-verified'])} | Gold and Silver, health relevance 2–3. Licence (`license_confirmed`, `license_class`), `public_access`, `latest_year`, `finest_unit_confirmed`, `join_key_confirmed` and the `evidence` URLs |",
      f"| `2-to-investigate.csv` | {len(inv)} | {len(old['2-to-investigate'])} | Bronze and Browser needed, health relevance 2–3; `what_to_check` says what is missing |",
      f"| `3-reference-layers.csv` | {len(ref)} | {len(old['3-reference-layers'])} | Boundaries, elevation, hydrography (health relevance 1 or unknown), with their tier |",
      f"| `4-trimmed.csv` | {len(T) + len(fail)} | {len(old['4-trimmed'])} | The trimmed rows plus {len(fail)} that failed verification and {len(DUPE_DROP_KEPT)} dropped as duplicates, each with the reason |", "",
      "### Movements between files", "", "| from | to | datasets |", "|---|---|---|"]
NAMES = {"1": "verified", "2": "to investigate", "3": "reference layers", "4": "trimmed or failed", "new": "new row"}
for (a_, b_), n in sorted(moved.items()):
    IS.append(f"| {NAMES[a_]} | {NAMES[b_]} | {n} |")
if not moved:
    IS.append("| (none) | | |")
IS += ["", "### Newly read", "",
       f"- {len(newly_read)} datasets sent to the content check on {DATE}; {len(newly_read_ok)} read, {len(newly_blocked)} blocked (JavaScript shells, challenge pages or empty responses; now in the browser queue).",
       f"- Of the {len(newly_read_ok)} read: " + ", ".join(f"{t} {n}" for t, n in Counter(r['verify_tier'] for r in newly_read_ok).most_common()) + ".",
       f"- Licences read on the page now cover {sum(1 for r in K if r['license_confirmed'])} datasets: " + ", ".join(f"{k} {v}" for k, v in lic.most_common()) + ". A `fail` on licence means no licence statement was found on the pages read, not that the data are restricted.",
       f"- Spatial unit confirmed for {sum(1 for r in K if r['chk_unit']=='pass')} datasets, corrected for {sum(1 for r in K if r['chk_unit']=='warn')}; years confirmed for {sum(1 for r in K if r['chk_temporal']=='pass')}; {sum(1 for r in K if r['chk_recency']=='fail')} found discontinued or pre-2000 (moved to Failed).",
       f"- `public_access` flags {len(reg)} registered or restricted datasets in the verified/investigate files.", "",
       "### Duplicates resolved", "",
       f"- {len(DUPES)} groups reviewed (`dupe-review-{DATE}.csv`): " + ", ".join(f"{k} {v}" for k, v in dup_dec.most_common()) + f". Only {sum(1 for d in DUPES if d['group'].startswith('D'))} of the 174 groups flagged in the harvest touch a kept row; the rest were trimmed or stayed as leads. {len(DUPE_DROP_KEPT)} kept rows dropped as duplicates; the alternates' URLs were recorded on the primary where the alternate was already outside the kept set.", "",
       "### Browser sample", "",
       f"- A person opened 25 Browser-needed rows in a real browser (`browser-sample-{DATE}.md`): 23 visited, 18 opened fine, 1 bot challenge, 1 sign-in or gone, 1 not found, 2 hosts unreachable. The suggested statuses were applied: " + ", ".join(f"{t} {n}" for t, n in Counter(BROWSER_TIER.values()).most_common()) + ". About 4 in 5 blocked pages open for a person, so 'Browser needed' mostly means bot-blocking, not dead links. The AHRQ SDOH Database row now points at its successor, the Community-Level Health (CLH) Database.", "",
       "### Mirrors' licences", ""]
for mid, (res, lic_, cls, det, ev_) in MIRROR_LICENSE.items():
    IS.append(f"- `{mid}`: {('**' + lic_ + '**') if lic_ else 'no licence stated'} ({res}); {det.split(';')[0]}.")
IS += ["", f"- Browser queue: {len(BQ)} rows on {len({x['domain'] for x in BQ})} domains (`browser-queue-{DATE}.csv` in the workspace).", "",
       f"The workbook `place-data-inventory-{DATE}.xlsx` has the same content with a colour-coded Verification grid, live summary counts, a coverage-of-assessment block and a Duplicates resolved tab."]
open(f"{TK}/ISSUE-update-{DATE}.md", "w", encoding="utf-8").write("\n".join(IS) + "\n")

# ----------------------------------------------------------------------------- 5. workbook
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
A = Font(name="Arial", size=10); AB = Font(name="Arial", size=10, bold=True); H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="14213D"); LINK = Font(name="Arial", size=10, color="2F5F98", underline="single")
FILL = {"pass": "D8EBD6", "warn": "FCE6C9", "fail": "F6D3D0", "n.a.": "EEEEEE", "not_checked": "EEEEEE"}
TFILL = {"Gold": "F5E6A3", "Silver": "E3E6EA", "Bronze": "EBD9C8", "Browser needed": "DCE8F5", "Failed": "F6D3D0"}
wb = Workbook(); wb.remove(wb.active)
HR = {"3": "3 · Direct exposure/determinant", "2": "2 · Context/confounder", "1": "1 · Reference geography", "0": "0 · None", "unknown": "Unknown"}


def sheet(name, cols, data, freeze="C2"):
    ws = wb.create_sheet(name)
    for j, (h, f, w) in enumerate(cols, 1):
        c = ws.cell(1, j, h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center"); ws.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(data, 2):
        for j, (h, f, w) in enumerate(cols, 1):
            v = f(r)
            if isinstance(v, str) and v.isdigit() and h in ("Latest year", "Health relevance (0-3)", "Score"):
                v = int(v)
            c = ws.cell(i, j, v); c.font = A
            if isinstance(v, str) and v.startswith("http"):
                c.hyperlink = v; c.font = LINK
            if h.startswith("✓") and v in FILL:
                c.fill = PatternFill("solid", fgColor=FILL[v])
            if h == "Tier" and v in TFILL:
                c.fill = PatternFill("solid", fgColor=TFILL[v])
    ws.freeze_panes = freeze; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(data), 1) + 1}"; ws.row_dimensions[1].height = 32
    return ws


# About
ab = wb.create_sheet("About"); ab.column_dimensions["A"].width = 26; ab.column_dimensions["B"].width = 120
ABOUT = [("Place-based data inventory: refined set", ""), ("Date", "28 October 2026 · OHDSI GIS Working Group · built 20–28 October 2026 (first-round assessment closed)"),
         ("What this is", f"{len(K)} public, place-based datasets that could be linked to where patients live, each with a verification tier. {tiers.get('Gold',0)+tiers.get('Silver',0)} are verified (Gold or Silver): their licence, years, spatial unit and access were read from the publisher's page, with the page URL kept as evidence."),
         ("How it was built", "1 Harvest: AI-assisted web sweeps plus scripted catalog-API listings (data.gov, HDX, Earthdata CMR, EU and national portals): 3,337 listings. 2 De-duplication: 3,160. 3 Split: 594 curated rows kept apart from 2,576 catalog leads; leads whose title names a spatial unit were added: 1,018. 4 Score: health relevance (0–3) and Gaia fit (joins today / needs a count or distance step / needs raster support / not linkable) from each listing. 5 Trim: 570 kept (relevant, data from 2000 on, not a one-off event, not hyper-local, not a portal). 6 URL health (21–22 Oct): scripted checks of every landing and access URL. 7 Content checks (22 and 28 Oct): an agent read the landing page and one documentation page of each live dataset (shortlist and step-3 groups on 22 Oct; the remaining live rows on 28 Oct). 8 Mirrors: four community mirrors or successors of withdrawn federal tools were added. 9 Duplicates (28 Oct): every flagged group among the kept rows was decided (see the Duplicates resolved tab)."),
         ("What 'verified' means", "Each check is pass / warn / fail / n.a. / not checked, with a detail and (for content checks) an evidence URL. Link checks: link live (2xx), redirect (same domain), TLS, soft-404 (page names the dataset), withdrawn (no retirement notice), access URL, format match, machine-readable metadata, persistent id, Wayback snapshot. Content checks: licence (+ class: open / attribution / non-commercial / registration / restricted), years (+ latest year), recency (pass 2020+, warn 2000–2019, fail before 2000 or discontinued), update cadence, version, finest unit, coverage, access route, join key, Gaia-loadable format. Workspace checks: in gaiaCatalog, recipe health from the deployment logs."),
         ("Tiers", "Gold: link, page, withdrawn, licence, unit, years and access all pass; recency pass or warn. Silver: link passes and no fail among licence, unit, years, recency. Bronze: link live but content not read or incomplete. Browser needed: the site blocks scripts or has a certificate problem; a person must open it. Failed: link dead, page generic, publisher says withdrawn, or data before 2000."),
         ("Score", "0–100, weighted: link 20, soft-404 10, withdrawn 10, licence 15, unit 15, years 7.5, recency 7.5, access 10, join key 10, machine metadata 5; pass = 1, warn = 0.5, fail = 0; normalised over the checks that apply."),
         ("Assumptions", "Links were checked automatically on 21–22 October 2026 and pages read on 22 and 28 October; a 'fail' on licence means no licence statement was found on the pages read, not that the data are restricted. Relevance and Gaia-fit are rule-based from each listing. Recency cut-off is 2000, when most medical data starts. Sites that block scripts were not worked around. No dataset was downloaded. Catalog leads carry a unit inferred from their title."),
         ("Tabs", "Summary (live formulas over the Verification tab, plus the coverage of the assessment) · Refined set (Gold and Silver) · Verification (every kept row, colour-coded checks) · Browser queue · Duplicates resolved · Trimmed (what was left out and why)."),
         ("Files", f"inventory-refined-{DATE}.csv (every column), inventory.csv (promoted Gold and Silver), browser-queue-{DATE}.csv, dupe-review-{DATE}.csv, close-report-{DATE}.md, ticket-{DATE}/ (GitHub issue package), scripts/verify_links.py and scripts/refine_2026_10_28.py in docs/discovery of the Gaia workspace.")]
for i, (k, v) in enumerate(ABOUT, 1):
    c = ab.cell(i, 1, k); c.font = Font(name="Arial", size=14 if i == 1 else 10, bold=True); c.alignment = Alignment(vertical="top")
    c = ab.cell(i, 2, v); c.font = A; c.alignment = Alignment(wrap_text=True, vertical="top")

# Verification (all keep=Y)
VER = [("Dataset", lambda r: r["dataset_name"], 50), ("Tier", lambda r: r["verify_tier"], 15), ("Score", lambda r: r["health_score"], 7), ("Theme", lambda r: r["theme"], 15),
       ("URL health", lambda r: r["url_health"], 11), ("License class", lambda r: r["license_class"] or ("not checked" if not r["partB_source"] else "unknown"), 13),
       ("Finest unit", lambda r: r.get("finest_unit_confirmed") or r["finest_unit"], 13), ("Gaia fit", lambda r: GF.get(r["gaia_fit"], r["gaia_fit"]), 22),
       ("Health relevance (0-3)", lambda r: r["health_relevance"], 10), ("In gaiaCatalog", lambda r: r["in_gaia"], 22), ("Source", lambda r: "Mirror" if r["mirror_of"] else ("Curated" if r["tier"] == "Curated" else "Catalog lead"), 10),
       ("Latest year", lambda r: r["latest_year_final"], 9), ("Public access", lambda r: r["public_access"], 11), ("Verified on", lambda r: r["verified_on"], 11)] + \
      [("✓ " + c[4:].replace("_", " "), (lambda c: lambda r: r[c])(c), 11) for c in ALL_CHECKS] + \
      [("Blockers", lambda r: r["verify_blockers"], 40), ("Landing page", lambda r: r["site_url"], 45), ("id", lambda r: r["id"], 30)]
sheet("Verification", VER, K)
vcol = {h: get_column_letter(j) for j, (h, _, _) in enumerate(VER, 1)}
n = len(K) + 1
rg = lambda h: f"Verification!${vcol[h]}$2:${vcol[h]}${n}"

# Refined set (Gold + Silver)
GS = [r for r in K if r["verify_tier"] in ("Gold", "Silver")]
GS.sort(key=lambda r: (r["theme"], -hrn(r), r["dataset_name"].lower()))
REF = [("Dataset", lambda r: r["dataset_name"], 52), ("Publisher", lambda r: r["organization"], 30), ("Theme", lambda r: r["theme"], 15), ("Coverage", lambda r: r.get("coverage_confirmed") or r["coverage"], 22),
       ("Finest unit", lambda r: r.get("finest_unit_confirmed") or r["finest_unit"], 13), ("Years", lambda r: r.get("temporal_confirmed") or r["temporal"], 28), ("Latest year", lambda r: r["latest_year_final"], 9),
       ("License class", lambda r: r["license_class"], 12), ("Licence (confirmed)", lambda r: r["license_confirmed"], 30), ("Access", lambda r: r.get("access_confirmed") or r["access_type"], 13),
       ("Format", lambda r: r.get("formats_confirmed") or r["format"], 28), ("Gaia fit", lambda r: GF.get(r["gaia_fit"], r["gaia_fit"]), 22), ("Health relevance (0-3)", lambda r: r["health_relevance"], 10),
       ("In gaiaCatalog", lambda r: r["in_gaia"], 22), ("Tier", lambda r: r["verify_tier"], 9), ("Score", lambda r: r["health_score"], 7), ("Landing page", lambda r: r["site_url"], 45), ("id", lambda r: r["id"], 28)]
sheet("Refined set", REF, GS)

# Browser queue
sheet("Browser queue", [("Domain", lambda x: x["domain"], 24), ("Dataset", lambda x: x["dataset_name"], 48), ("Reason", lambda x: x["reason"], 30), ("What to confirm", lambda x: x["what_to_confirm"], 70),
                        ("Shortlist", lambda x: "yes" if x["priority"] else "", 8), ("Landing page", lambda x: x["site_url"], 45), ("Archive copy", lambda x: x["archive_copy"], 35), ("id", lambda x: x["id"], 28)], BQ, freeze="B2")
# Duplicates resolved
sheet("Duplicates resolved", [("Group", lambda d: d["group"], 9), ("Datasets (ids, primary first)", lambda d: d["ids"].replace(";", "; "), 60), ("Decision", lambda d: d["decision"], 16),
                               ("Reason", lambda d: d["reason"], 90), ("Applied", lambda d: d["applied"], 60)], DUPES, freeze="B2")
# Trimmed
TR = sorted(T, key=lambda r: (r["trim_reason"], r["dataset_name"].lower())) + fail
sheet("Trimmed", [("Dataset", lambda r: r["dataset_name"], 55), ("Why left out", why, 70), ("Publisher", lambda r: r["organization"], 28), ("Latest year", lambda r: r.get("latest_year_final") or r.get("latest_year", ""), 9),
                  ("Source", lambda r: "Curated" if r["tier"] == "Curated" else "Catalog lead", 11), ("Landing page", lambda r: r["site_url"], 45), ("id", lambda r: r["id"], 28)], TR, freeze="B2")

# Summary (formulas only, except the funnel)
s = wb.create_sheet("Summary", 0); s.column_dimensions["A"].width = 44
for Lc in "BCDEFGH": s.column_dimensions[Lc].width = 15
FORMULAS = []  # (cell, formula, python expected) for the check


def hdr(row, vals):
    for j, v in enumerate(vals, 1):
        c = s.cell(row, j, v); c.font = H; c.fill = HF


def put(row, col, formula, expected):
    c = s.cell(row, col, formula); c.font = A; FORMULAS.append((f"{get_column_letter(col)}{row}", formula, expected))


s["A1"] = "Place-based data inventory · refined set · summary"; s["A1"].font = Font(name="Arial", size=14, bold=True)
s["A2"] = "Counts are live formulas over the Verification tab (every kept dataset), except the funnel (fixed counts; sources named)."; s["A2"].font = Font(name="Arial", size=10, italic=True)
TIERS = ["Gold", "Silver", "Bronze", "Browser needed", "Failed"]
themes = sorted({r["theme"] for r in K})
row = 4; hdr(row, ["Tier × theme"] + TIERS + ["Total"])
for i, t in enumerate(themes):
    rr = row + 1 + i; s.cell(rr, 1, t).font = A
    for j, ti in enumerate(TIERS):
        put(rr, 2 + j, f'=COUNTIFS({rg("Theme")},$A{rr},{rg("Tier")},"{ti}")', sum(1 for r in K if r["theme"] == t and r["verify_tier"] == ti))
    put(rr, 7, f"=SUM(B{rr}:F{rr})", sum(1 for r in K if r["theme"] == t))
rr = row + 1 + len(themes); s.cell(rr, 1, "Total").font = AB
TIER_GS_REF = f"=B{rr}+C{rr}"  # Gold + Silver totals (the 2026-10-23 workbook pointed this at the wrong cell)
for j in range(2, 8):
    put(rr, j, f"=SUM({get_column_letter(j)}{row+1}:{get_column_letter(j)}{rr-1})", sum(1 for r in K if (j == 7 or r["verify_tier"] == TIERS[j-2])))
    s.cell(rr, j).font = AB
row = rr + 2; hdr(row, ["URL health", "Datasets"])
for i, u in enumerate(["healthy", "degraded", "moved", "blocked", "dead", "n.a."]):
    s.cell(row + 1 + i, 1, u).font = A; put(row + 1 + i, 2, f'=COUNTIF({rg("URL health")},A{row+1+i})', sum(1 for r in K if r["url_health"] == u))
row = row + 8; hdr(row, ["License class (content checks)", "Datasets"])
LC = ["open", "attribution", "non-commercial", "registration", "restricted", "unknown", "not checked"]
for i, u in enumerate(LC):
    s.cell(row + 1 + i, 1, u).font = A
    put(row + 1 + i, 2, f'=COUNTIF({rg("License class")},A{row+1+i})', sum(1 for r in K if (r["license_class"] or ("not checked" if not r["partB_source"] else "unknown")) == u))
row = row + len(LC) + 2; hdr(row, ["Finest unit", "Datasets"])
units = [u for u, _ in Counter((r.get("finest_unit_confirmed") or r["finest_unit"]) for r in K).most_common()]
for i, u in enumerate(units):
    s.cell(row + 1 + i, 1, u or "(blank)").font = A
    put(row + 1 + i, 2, f'=COUNTIF({rg("Finest unit")},A{row+1+i})' if u else f'=COUNTBLANK({rg("Finest unit")})', sum(1 for r in K if (r.get("finest_unit_confirmed") or r["finest_unit"]) == u))
row = row + len(units) + 2; hdr(row, ["Gaia fit × health relevance", "Health 3", "Health 2", "Health 1", "Health 0", "Unknown", "Total"])
for i, g in enumerate(GF.values()):
    rr = row + 1 + i; s.cell(rr, 1, g).font = A
    for j, h in enumerate(["3", "2", "1", "0", "unknown"]):
        put(rr, 2 + j, f'=COUNTIFS({rg("Gaia fit")},$A{rr},{rg("Health relevance (0-3)")},{h if h.isdigit() else chr(34)+h+chr(34)})',
            sum(1 for r in K if GF.get(r["gaia_fit"], r["gaia_fit"]) == g and r["health_relevance"] == h))
    put(rr, 7, f'=COUNTIF({rg("Gaia fit")},$A{rr})', sum(1 for r in K if GF.get(r["gaia_fit"], r["gaia_fit"]) == g))
row = row + len(GF) + 2; hdr(row, ["Already in gaiaCatalog", "Datasets"])
s.cell(row + 1, 1, "Kept datasets with a gaiaCatalog folder").font = A; put(row + 1, 2, f'=COUNTIF({rg("In gaiaCatalog")},"?*")', sum(1 for r in K if r["in_gaia"].strip()))
s.cell(row + 2, 1, "…of which Gold or Silver").font = A
put(row + 2, 2, f'=COUNTIFS({rg("In gaiaCatalog")},"?*",{rg("Tier")},"Gold")+COUNTIFS({rg("In gaiaCatalog")},"?*",{rg("Tier")},"Silver")', sum(1 for r in K if r["in_gaia"].strip() and r["verify_tier"] in ("Gold", "Silver")))
s.cell(row + 3, 1, "Public access: registration or restricted").font = A
put(row + 3, 2, f'=COUNTIF({rg("Public access")},"registration")+COUNTIF({rg("Public access")},"restricted")', sum(1 for r in K if r["public_access"] in ("registration", "restricted")))
row = row + 5; hdr(row, ["Coverage of assessment", "Datasets", "Of which Gold/Silver"])
COV = [("Read by part B (content checks, 22 or 28 Oct)", lambda r: bool(r["partB_source"]) and r["partB_blocked"] != "TRUE"),
       ("Sent to part B but blocked at the content step", lambda r: r["partB_blocked"] == "TRUE"),
       ("Not read: browser queue (site blocks scripts or certificate problem)", lambda r: not r["partB_source"] and r["url_health"] == "blocked" and r["browser_tier"] in ("", "Browser needed")),
       ("Not read: link dead or withdrawn (Failed)", lambda r: not r["partB_source"] and r["verify_tier"] == "Failed"),
       ("Not read: no URL", lambda r: not r["partB_source"] and r["url_health"] == "n.a."),
       ("Checked in a browser by a person (sample, 28 Oct)", lambda r: bool(r["browser_tier"]) and r["browser_tier"] != "Browser needed"),
       ("Not read: other", lambda r: not r["partB_source"] and not r["browser_tier"] and r["url_health"] not in ("blocked", "n.a.") and r["verify_tier"] != "Failed")]
for i, (lab, fn) in enumerate(COV):
    rr = row + 1 + i; s.cell(rr, 1, lab).font = A
    s.cell(rr, 2, sum(1 for r in K if fn(r))).font = Font(name="Arial", size=10, color="0000FF")
    s.cell(rr, 3, sum(1 for r in K if fn(r) and r["verify_tier"] in ("Gold", "Silver"))).font = Font(name="Arial", size=10, color="0000FF")
rr = row + 1 + len(COV); s.cell(rr, 1, "Total kept").font = AB; s.cell(rr, 2, len(K)).font = AB
row = rr + 2; hdr(row, ["How the list was built (fixed counts)", "Rows", "Source"])
FUN = [("Harvested: agent web sweeps + catalog-API listings", 3337, "harvest-2026-10-20-summary.md"), ("After removing duplicates", 3160, "harvest-2026-10-20-dedupe-report.md"),
       ("Curated rows, or a catalog lead with a stated spatial unit", 1018, "next-pass-2026-10-20-report.md; place-data-inventory-2026-10-21.xlsx"),
       ("Kept after trimming", 570, "trim-2026-10-22-report.md"), ("Mirrors / successors of withdrawn tools added", len(mirrors), f"refine-report-{PREV}.md"), ("Dropped as duplicates (28 Oct)", len(DUPE_DROP_KEPT), f"dupe-review-{DATE}.csv"), ("Trimmed after a browser check (28 Oct)", sum(1 for r in T if r["trim_reason"].startswith(BROWSER_NOTE)), f"browser-sample-{DATE}.md"),
       ("Read for content (part B, 22 and 28 Oct)", sum(1 for r in K if r["partB_source"]), f"close-report-{DATE}.md"), ("Verified (Gold + Silver)", TIER_GS_REF, "Summary, tier totals")]
for i, (a, b, c) in enumerate(FUN):
    rr = row + 1 + i; s.cell(rr, 1, a).font = A; s.cell(rr, 2, b).font = Font(name="Arial", size=10, color="0000FF" if isinstance(b, int) else "000000"); s.cell(rr, 3, c).font = A
s.cell(row + len(FUN) + 1, 1, "Blue numbers are fixed counts from the files named in column C (docs/discovery, 20–28 October 2026).").font = Font(name="Arial", size=9, italic=True)
s.column_dimensions["C"].width = 44
for ws in wb.worksheets:
    for rowc in ws.iter_rows():
        for c in rowc:
            if c.value is not None and (c.font is None or c.font.name != "Arial"):
                c.font = A
wb.calculation.fullCalcOnLoad = True
XL = f"{D}/place-data-inventory-{DATE}.xlsx"
wb.save(XL)

# formula check: evaluate every Summary formula against the sheet values with a small evaluator, compare with the Python counts
import openpyxl
wb2 = openpyxl.load_workbook(XL, data_only=False)


def col_vals(ref):
    m = re.match(r"(\w[\w ]*)!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", ref)
    ws = wb2[m.group(1)]
    return [ws[f"{m.group(2)}{i}"].value for i in range(int(m.group(3)), int(m.group(5)) + 1)]


def crit(v, c):
    c = c.strip()
    if c.startswith('"') and c.endswith('"'):
        c = c[1:-1]
        if c == "?*":
            return isinstance(v, str) and len(v) >= 1
        return str(v if v is not None else "") == c
    if re.fullmatch(r"-?\d+", c):
        return (str(v) == c) or (v == int(c))
    if re.fullmatch(r"\$?[A-Z]+\d+", c):
        return str(v if v is not None else "") == str(wb2["Summary"][c.replace("$", "")].value if wb2["Summary"][c.replace("$", "")].value is not None else "")
    raise ValueError(c)


def ev(formula, cache):
    f = formula[1:]
    total = 0
    for term in re.split(r"\+(?=COUNTIF)", f):
        m = re.fullmatch(r"COUNTIF\((.+?),(.+)\)", term)
        if m:
            vals = col_vals(m.group(1)); total += sum(1 for v in vals if crit(v, m.group(2))); continue
        m = re.fullmatch(r"COUNTIFS\((.+)\)", term)
        if m:
            parts = re.split(r",(?=(?:[^\"]*\"[^\"]*\")*[^\"]*$)", m.group(1)); cols = [col_vals(parts[i]) for i in range(0, len(parts), 2)]
            crits = [parts[i] for i in range(1, len(parts), 2)]
            total += sum(1 for k in range(len(cols[0])) if all(crit(cols[j][k], crits[j]) for j in range(len(cols)))); continue
        m = re.fullmatch(r"COUNTBLANK\((.+)\)", term)
        if m:
            total += sum(1 for v in col_vals(m.group(1)) if v in (None, "")); continue
        m = re.fullmatch(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", term)
        if m:
            for i in range(int(m.group(2)), int(m.group(4)) + 1):
                for cl in range(openpyxl.utils.column_index_from_string(m.group(1)), openpyxl.utils.column_index_from_string(m.group(3)) + 1):
                    cell = f"{get_column_letter(cl)}{i}"; v = wb2["Summary"][cell].value
                    total += cache[cell] if isinstance(v, str) and v.startswith("=") else (v or 0)
            continue
        raise ValueError(term)
    return total


cache = {}; mism = []
for cell, formula, expected in FORMULAS:
    try:
        got = ev(formula, cache); cache[cell] = got
        if got != expected:
            mism.append((cell, formula, expected, got))
    except Exception as e:
        mism.append((cell, formula, expected, f"ERROR {e}"))
formula_note = f"{len(FORMULAS)} Summary formulas evaluated; {len(mism)} mismatches or errors" + (": " + str(mism[:3]) if mism else "")
print(formula_note, file=sys.stderr)

# ----------------------------------------------------------------------------- 6. close report
def ctab(checks):
    out = ["| Check | pass | warn | fail | n.a. | not_checked |", "|---|---|---|---|---|---|"]
    for c in checks:
        cnt = Counter(r[c] for r in K)
        out.append(f"| `{c}` | {cnt.get('pass',0)} | {cnt.get('warn',0)} | {cnt.get('fail',0)} | {cnt.get('n.a.',0)} | {cnt.get('not_checked',0)} |")
    return out


nB = sum(1 for r in K if r["partB_source"]); nBb = sum(1 for r in K if r["partB_blocked"] == "TRUE")
files_D = sorted(glob.glob(f"{D}/partB/results/D*.json"))
USAGE = json.load(open(f"{D}/partB/usage-close.json")) if os.path.exists(f"{D}/partB/usage-close.json") else {}
_mt = [os.path.getmtime(f) for f in files_D]
t_first = datetime.datetime.fromtimestamp(min(_mt), datetime.timezone.utc).strftime("%Y-%m-%d %H:%M") if _mt else "-"
t_last = datetime.datetime.fromtimestamp(max(_mt), datetime.timezone.utc).strftime("%Y-%m-%d %H:%M") if _mt else "-"
gs_theme = Counter(r["theme"] for r in GS)
queue_done = sum(1 for i in QUEUE_CLOSE if i in partB); queue_left = [i for i in QUEUE_CLOSE if i not in partB]
R = [f"# Close report: first-round assessment of the place-based data inventory ({DATE})", "",
     f"**This assumes:** the 2026-10-23 refined set (574 kept rows) is the scope; part A results are unchanged (2026-10-21 run plus the 2026-10-22 re-check); part B is an agent reading the landing page and at most one documentation page per dataset (no downloads, no bot-protection workarounds, no guessing: every confirmed value has an evidence URL); recency cut-off 2000 (Robert, 2026-10-22). Duplicate decisions were made from the listings (title, publisher, URL, coverage, unit, years), without re-reading pages.", "",
     "## Funnel", "", "| Step | Datasets |", "|---|---|", "| Harvested | 3,337 |", "| After de-duplication | 3,160 |", "| Curated, or a lead with a stated unit | 1,018 |", "| Kept after trimming | 570 |",
     f"| + mirrors / successors of withdrawn tools | {len(mirrors)} |", f"| − dropped as duplicates ({DATE}) | {len(DUPE_DROP_KEPT)} |", f"| − trimmed after a browser check ({DATE}) | {sum(1 for r in T if r['trim_reason'].startswith(BROWSER_NOTE))} |", f"| Kept | {len(K)} |",
     f"| **Read by part B: {nB} of {len(PREV_REFINED)}** (22 Oct: {sum(1 for r in K if r['partB_source'] and not r['partB_source'].startswith('D'))}; 28 Oct: {sum(1 for r in K if r['partB_source'].startswith('D'))}; blocked at the content step: {nBb}) | {nB} |",
     f"| Checked in a browser by a person (sample of the browser queue, {DATE}) | {sum(1 for r in K if r['browser_tier'] and r['browser_tier'] != 'Browser needed')} |",
     f"| Not read: browser queue {sum(1 for r in K if not r['partB_source'] and r['url_health']=='blocked')}, dead {sum(1 for r in K if not r['partB_source'] and r['url_health']=='dead' and not r['browser_tier'])}, no URL {sum(1 for r in K if not r['partB_source'] and r['url_health']=='n.a.')}, left in the close queue {len(queue_left)} | {len(K) - nB - sum(1 for r in K if r['browser_tier'] and r['browser_tier'] != 'Browser needed')} |",
     f"| **Verified: Gold {tiers.get('Gold',0)} + Silver {tiers.get('Silver',0)}** | **{tiers.get('Gold',0)+tiers.get('Silver',0)}** |", "",
     "## Tiers before and after", "", "| Tier | 2026-10-23 | 2026-10-28 | of which shortlist |", "|---|---|---|---|"] + \
    [f"| {t} | {tiers_prev.get(t,0)} | {tiers.get(t,0)} | {sum(1 for r in K if r['verify_tier']==t and r['priority'])} |" for t in TIERS] + ["",
     "Movements: " + (", ".join(f"{a} → {b}: {n}" for (a, b), n in sorted(tier_moves.items(), key=lambda x: -x[1])) or "none") + ".", "",
     "## Step 1: part B on the remaining live rows", "",
     f"Queue `partB/queue-close.json`: {len(QUEUE_CLOSE)} Bronze rows never read (the one Bronze row without a URL was excluded), ordered health 3 with unknown/non-linkable unit ({sum(1 for i in QUEUE_CLOSE if PREV_REFINED[i]['health_relevance']=='3' and PREV_REFINED[i]['gaia_fit'] in ('unknown', 'not_linkable'))}), health 2 needing a count/distance step ({sum(1 for i in QUEUE_CLOSE if PREV_REFINED[i]['health_relevance']=='2' and PREV_REFINED[i]['gaia_fit']=='needs_aggregation')}), health 2 needing raster ({sum(1 for i in QUEUE_CLOSE if PREV_REFINED[i]['health_relevance']=='2' and PREV_REFINED[i]['gaia_fit']=='needs_raster')}), reference layers ({sum(1 for i in QUEUE_CLOSE if PREV_REFINED[i]['health_relevance']=='1')}), the rest. "
     f"Done: {queue_done} of {len(QUEUE_CLOSE)} in {len(files_D)} agent runs ({len({os.path.basename(f).split('-')[0] for f in files_D})} batches); {len(queue_left)} left" + (": " + ", ".join(queue_left[:12]) + ("…" if len(queue_left) > 12 else "") if queue_left else "") + ".", "",
     f"Results on the newly read rows: " + ", ".join(f"{t} {n}" for t, n in Counter(r['verify_tier'] for r in newly_read).most_common()) + f" (blocked at the content step: {len(newly_blocked)}, now Browser needed).", "",
     "Part B (content, all kept rows, 22 and 28 Oct):", ""] + ctab(B_CHECKS) + ["",
     "## Gold and Silver by theme", "", "| Theme | Gold | Silver | Total |", "|---|---|---|---|"] + \
    [f"| {t} | {sum(1 for r in GS if r['theme']==t and r['verify_tier']=='Gold')} | {sum(1 for r in GS if r['theme']==t and r['verify_tier']=='Silver')} | {gs_theme[t]} |" for t in sorted(gs_theme)] + ["",
     "## Step 2: duplicate decisions", "",
     f"Of the 174 groups flagged in the harvest, only {sum(1 for d in DUPES if d['group'].startswith('D'))} touch a kept row ({sum(1 for d in DUPES if d['group'].startswith('D') and len([i for i in d['ids'].split(';') if i in PREV_REFINED]) >= 2)} with two or more kept members); the rest were trimmed on 2026-10-22 or stayed as leads. A scan of the kept set for rows sharing a landing page added {sum(1 for d in DUPES if d['group'].startswith('X'))} groups (X-numbered). Decisions (`dupe-review-{DATE}.csv`): " + ", ".join(f"{k} {v}" for k, v in dup_dec.most_common()) + f". Kept rows dropped: {len(DUPE_DROP_KEPT)}. Where the alternate was already outside the kept set, its URL was recorded on the primary (notes, or `access_url` when empty: {', '.join(DUPE_ACCESS) or 'none'}).", "",
     "| group | decision | ids |", "|---|---|---|"] + [f"| {d['group']} | {d['decision']} | {d['ids'].replace(';', '; ')} |" for d in DUPES if d['decision'] != 'distinct products'] + [
     f"| {sum(1 for d in DUPES if d['decision']=='distinct products')} groups | distinct products | flag cleared |", "",
     "## Browser sample (a person, 2026-10-28)", "",
     f"25 Browser-needed rows were opened in a real browser (`browser-sample-{DATE}.md`, `.csv`): 23 visited, 18 opened. The suggested status changes were applied: " + ", ".join(f"{t} {n}" for t, n in Counter(BROWSER_TIER.values()).most_common()) + f". Fields read in the browser (licence, years, unit, access, formats) are in the refined CSV with the note '{BROWSER_NOTE}'; `browser_tier` records the applied status. `us_ahrq_sdoh_database` was repointed to the AHRQ Community-Level Health (CLH) Database page. NRDC stays Browser needed (Cloudflare challenge).", "",
     "## Step 3: loose ends", "", "Mirror licences (pages read 2026-10-28):", "", "| row | licence | result | where looked |", "|---|---|---|---|"] + \
    [f"| `{m}` | {lic_ or 'none stated'} | {res} | {ev_.replace(';', ';<br>')} |" for m, (res, lic_, cls, det, ev_) in MIRROR_LICENSE.items()] + ["",
     "Suggested question to PEDP (no general contact address on the pages; the About page lists fundraising@publicenvirodata.org, and the CEJST and HIFLD Next pages have feedback forms): *\"Under what licence or terms do you make the EJScreen reconstruction (screening-tools.com/epa-ejscreen and the Harvard Dataverse deposit), the CEJST v2.0 copy and HIFLD Next available for reuse? The federal originals were US Government works in the public domain; could you state that, or your own terms, on each page or in the repository?\"*", "",
     f"Notes for Tim updated: `notes-for-tim-{DATE}.md` (draft, not sent). Decision memo: `decisions-needed-{DATE}.md`.", "",
     "## Choices made (decided without asking)", "",
     "- The close queue excludes the one Bronze row without a URL (`us_mdc_2024_dormir_study`); dead and blocked rows were already outside Bronze.",
     "- Batches of 25 were run with 5 agents each as before, but consecutive batches overlapped (the next batch started once most agents of the previous one had returned) to fit the time-box.",
     "- Part B was stopped on Robert's instruction after a usage limit cut batches D07 and D08 short: valid result files were kept, agents still running were stopped, and no further batches were launched. D07-a2's file was invalid when first checked and valid after the agent rewrote it, so it is included.",
     "- The browser sample's suggested statuses are applied as tier overrides (`browser_tier`); the check columns record what the page showed (for example 'public domain, assumed' is a licence warn, not a pass), so a Gold here can carry a licence warn.",
     "- Rows that the content step found to be JavaScript-only shells or challenge pages are `blocked` and move to Browser needed, as on 2026-10-23.",
     "- Duplicate groups where only one member survived the trim are recorded as decided (the alternate's URL noted on the primary) rather than ignored; no kept row was dropped, because every pair of kept rows turned out to be distinct products.",
     "- The six X-groups from the same-landing-page scan are an addition to the brief; all were judged distinct, so nothing changed.",
     "- The USGS 3DHP mirror's licence was filled from the publisher-level USGS Copyrights and Credits page, as the brief asked for a terms page; this is the only row whose licence comes from a publisher-level page. The same rule has not been applied to the 100-odd US federal rows whose pages state nothing (still Robert's call; see the decision memo and next steps).",
     "- The three PEDP mirrors stay at `chk_license` fail (nothing stated), with the conflicting code licences on the CEJST repository recorded in the detail.",
     "- `verified_on` for rows read on 28 Oct is 2026-10-28; rows read on 22 Oct keep 2026-10-22.",
     "- The ticket's `1-verified.csv` gained `finest_unit_confirmed`, `join_key_confirmed` and `evidence` columns; `link_status_2026_10_22` was renamed `link_status`.",
     "- The workbook's funnel line 'Verified (Gold + Silver)' now points at the tier totals; in the 2026-10-23 workbook it referenced a Gaia-fit count cell (a bug in the 10-23 script; the CSVs and reports were unaffected).",
     "- Workbook recalculation: LibreOffice is not installed; every Summary formula was evaluated by a small Python evaluator against the sheet values and compared with the Python counts (" + formula_note + "). Excel will recalculate on open (fullCalcOnLoad).", "",
     "## Time and tokens", "",
     f"- Part B close run: {len(files_D)} agent runs (5 rows each), {len(newly_read)} rows; subagent usage as reported by the harness: " + (f"{sum(USAGE.values())/1e6:.2f}M tokens (" + ", ".join(f"{k} {v/1e3:.0f}k" for k, v in sorted(USAGE.items())) + ")" if USAGE else "not recorded") + ", plus the coordinating session. Steps 2–4 used no LLM apart from the coordinator.",
     f"- Wall time: run started 2026-10-28 19:48 UTC; first D result written {t_first}, last {t_last} UTC; package written {datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC.", "",
     "## What remains", "",
     f"- Browser queue: {len(BQ)} rows on {len({x['domain'] for x in BQ})} domains (`browser-queue-{DATE}.csv`), including the {len(newly_blocked)} JavaScript-only pages found today. `browser-sample-prompt-{DATE}.md` has a sample.",
     "- Round-2 gap sweeps (`docs/discovery/round2/`, separate session): health-facility registries, transport, hazards, food access outside the US, sub-national health outcomes.",
     f"- Robert's decisions (`decisions-needed-{DATE}.md`): the two gaiaCatalog entries, the manuscript's '113 curated datasets' wording, the three PEDP mirrors without a licence; and the licence policy for US federal pages that state nothing ({sum(1 for r in K if r['chk_license']=='fail' and r['partB_blocked']!='TRUE')} rows at fail).",
     f"- {len(queue_left)} close-queue rows not read (if any): same method, `partB/batches/D*.json`.",
     "- Promotion of the Gold/Silver rows into the catalog pipeline; send the notes for Tim once read.",
     "- Monthly re-run of `verify_links.py`."]
open(f"{D}/close-report-{DATE}.md", "w", encoding="utf-8").write("\n".join(R) + "\n")
print(f"kept {len(K)} rows; tiers {dict(tiers)}; promoted {len(promoted)}; browser queue {len(BQ)}; ticket 1/2/3/4 = {len(ver)}/{len(inv)}/{len(ref)}/{len(T)+len(fail)}; queue done {queue_done}/{len(QUEUE_CLOSE)}", file=sys.stderr)
